"""
Модуль администратора: Управление Пользователями (Admin - Users).
Содержит логику создания, удаления пользователей, принудительного сброса паролей
и массового назначения отчетов конкретному пользователю.
"""
from flask import request, redirect, url_for, flash, send_file
from flask_login import login_required, current_user
import io
import openpyxl
from app.extensions import db, limiter
from app.admin import admin_bp
from app.models import User, ReportSubmission, ReportTemplate
from app.utils import log_action
from app.auth.decorators import roles_required

# ==========================================
# УПРАВЛЕНИЕ ПОЛЬЗОВАТЕЛЯМИ
# ==========================================

@admin_bp.route('/create_user', methods=['POST'])
@login_required
@roles_required('admin', 'manager')
def create_user():
    """
    Создание нового аккаунта. 
    Роль по умолчанию: 'user' (учреждение, сдающее отчет).
    Также может создавать роль 'manager' (ответственный - полный контроль отчетов).
    """
    username = request.form.get('username').strip()
    password = request.form.get('password')

    # Защита: роль только из белого списка, чтобы нельзя передать role=admin через POST
    ALLOWED_ROLES = ['user', 'manager']
    role = request.form.get('role', 'user')
    if role not in ALLOWED_ROLES:
        role = 'user'

    description = request.form.get('description', '').strip()
    group = request.form.get('group', None)
    if group == '':
        group = None

    if User.query.filter_by(username=username).first():
        flash('Пользователь с таким логином уже существует')
        return redirect(url_for('admin.dashboard', tab='usersTab'))

    user = User(username=username, role=role, description=description, group=group)
    user.set_password(password)
    
    # Если создает менеджер, автоматически привязываем пользователя к его отделу
    if current_user.role == 'manager':
        from app.utils import get_manager_department
        dept = get_manager_department(current_user)
        if dept:
            user.department_id = dept.id
    elif current_user.role == 'admin':
        dept_id_val = request.form.get('department_id')
        if dept_id_val and dept_id_val.isdigit():
            user.department_id = int(dept_id_val)
            
    db.session.add(user)
    db.session.commit()
    log_action('Создание пользователя', f'Создан новый пользователь: {username} с ролью {role}')
    return redirect(url_for('admin.dashboard', tab='usersTab'))

@admin_bp.route('/delete_user/<int:user_id>', methods=['POST'])
@login_required
@roles_required('admin', 'manager')
def delete_user(user_id):
    """
    Удаление пользователя и каскадное удаление всех его сданных отчетов.
    Вызывается из вкладки "Пользователи".
    """
    user = User.query.get_or_404(user_id)
    
    # Защита для менеджера
    if current_user.role == 'manager':
        from app.utils import get_manager_department
        dept = get_manager_department(current_user)
        if not dept or user.department_id != dept.id:
            flash('Доступ запрещен')
            return redirect(url_for('admin.dashboard', tab='usersTab'))
    # Очищаем связанные данные, чтобы не сломать внешние ключи
    ReportSubmission.query.filter_by(user_id=user.id).delete()
    username = user.username
    db.session.delete(user)
    db.session.commit()
    log_action('Удаление пользователя', f'Пользователь {username} и все его отчеты удалены')
    return redirect(url_for('admin.dashboard', tab='usersTab'))

@admin_bp.route('/edit_user/<int:user_id>', methods=['POST'])
@login_required
@roles_required('admin', 'manager')
def edit_user(user_id):
    """Единый маршрут для редактирования всех данных пользователя (основные данные, пароль, доступы)."""
    user = User.query.get_or_404(user_id)
    
    # Защита для менеджера
    if current_user.role == 'manager':
        from app.utils import get_manager_department
        dept = get_manager_department(current_user)
        if not dept or user.department_id != dept.id:
            flash('Доступ запрещен')
            return redirect(url_for('admin.dashboard', tab='usersTab'))
    
    # 1. Основные данные
    username = request.form.get('username')
    if username:
        # Проверяем уникальность логина, если он был изменен
        if username != user.username and User.query.filter_by(username=username).first():
            flash('Пользователь с таким логином уже существует')
            return redirect(url_for('admin.dashboard', tab='usersTab'))
        user.username = username
        
    user.description = request.form.get('description', '').strip()
    
    # Роль и группа
    role = request.form.get('role')
    if role in ['user', 'manager']:
        user.role = role
        
    group = request.form.get('group', None)
    if group == '':
        group = None
    user.group = group

    if current_user.role == 'admin':
        dept_id_val = request.form.get('department_id')
        user.department_id = int(dept_id_val) if (dept_id_val and dept_id_val.isdigit()) else None

    # 2. Сброс пароля (если заполнено поле)
    new_password = request.form.get('new_password')
    if new_password and new_password.strip():
        user.set_password(new_password.strip())
        
    # 3. Права доступа к отчетам
    user.assigned_templates = [] # Очищаем старые доступы
    for t_id in request.form.getlist('template_ids'):
        template = ReportTemplate.query.get(t_id)
        if template:
            user.assigned_templates.append(template)

    db.session.commit()
    log_action('Редактирование пользователя', f'Обновлен профиль пользователя {user.username}')
    return redirect(url_for('admin.dashboard', tab='usersTab'))

@admin_bp.route('/change_my_password', methods=['POST'])
@login_required
@limiter.limit("5 per minute")
def change_my_password():
    """Смена собственного пароля (вызывается через модальное окно профиля)."""
    new_password = request.form.get('new_password')
    if new_password:
        current_user.set_password(new_password)
        db.session.commit()
        log_action('Смена пароля', f'Пользователь {current_user.username} сменил свой пароль')
    # Защита от Open Redirect: возвращаемся только на внутренний маршрут 
    return redirect(url_for('admin.dashboard', tab='usersTab'))


@admin_bp.route('/users/download_template', methods=['GET'])
@login_required
@roles_required('admin')
def download_template():
    """Генерирует и скачивает шаблон Excel для массовой загрузки."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Пользователи"
    
    # Заголовки (точно как на скриншоте)
    headers = ["Логин", "Пароль", "Описание", "группа"]
    ws.append(headers)
    
    # Стилизация заголовков (полужирный)
    from openpyxl.styles import Font
    for cell in ws[1]:
        cell.font = Font(bold=True)
    
    # Настраиваем ширину колонок для красоты
    ws.column_dimensions['A'].width = 20
    ws.column_dimensions['B'].width = 20
    ws.column_dimensions['C'].width = 40
    ws.column_dimensions['D'].width = 20
    
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    
    return send_file(
        output,
        as_attachment=True,
        download_name="Шаблон_Пользователи.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

@admin_bp.route('/users/bulk_upload', methods=['POST'])
@login_required
@roles_required('admin')
def bulk_upload():
    """Обработка загруженного Excel-файла."""
    if 'file' not in request.files:
        flash("Файл не выбран")
        return redirect(url_for('admin.dashboard', tab='usersTab'))
        
    file = request.files['file']
    if file.filename == '':
        flash("Файл не выбран")
        return redirect(url_for('admin.dashboard', tab='usersTab'))
        
    role = request.form.get('role', 'user')
    if role not in ['user', 'manager']:
        role = 'user'
        
    try:
        wb = openpyxl.load_workbook(file)
        ws = wb.active
        
        added_count = 0
        skipped_count = 0
        
        # Читаем со второй строки (пропуская заголовки)
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or not row[0]: # Если строка пустая или нет логина
                continue
                
            username = str(row[0]).strip()
            password = str(row[1]).strip() if len(row) > 1 and row[1] is not None else ""
            description = str(row[2]).strip() if len(row) > 2 and row[2] is not None else ""
            group = str(row[3]).strip() if len(row) > 3 and row[3] is not None else None
            
            if group == '':
                group = None
                
            if not username or not password:
                skipped_count += 1
                continue
                
            # Проверка дубликата
            existing_user = User.query.filter_by(username=username).first()
            if existing_user:
                skipped_count += 1
                continue
                
            # Создание
            new_user = User(username=username, role=role, description=description, group=group)
            new_user.set_password(password)
            db.session.add(new_user)
            added_count += 1
            
        db.session.commit()
        log_action('Массовая загрузка', f'Добавлено пользователей: {added_count}, пропущено: {skipped_count}')
        flash(f"Успешно загружено: {added_count}. Пропущено (уже существуют или ошибка данных): {skipped_count}")
        
    except Exception as e:
        flash(f"Ошибка при обработке файла: {str(e)}")
        
    return redirect(url_for('admin.dashboard', tab='usersTab'))


@admin_bp.route('/users/bulk_delete', methods=['POST'])
@login_required
@roles_required('admin', 'manager')
def bulk_delete_users():
    """
    Массовое удаление выбранных пользователей.
    Каскадно удаляет связанные отчеты и черновики, очищает внешние ключи.
    """
    user_ids = request.form.getlist('user_ids')
    if not user_ids:
        flash('Не выбрано ни одного пользователя для удаления', 'warning')
        return redirect(url_for('admin.dashboard', tab='usersTab'))
        
    try:
        user_ids = [int(uid) for uid in user_ids if uid and str(uid).isdigit()]
    except (ValueError, TypeError):
        flash('Некорректный список идентификаторов', 'danger')
        return redirect(url_for('admin.dashboard', tab='usersTab'))
        
    if not user_ids:
        flash('Не выбрано ни одного пользователя', 'warning')
        return redirect(url_for('admin.dashboard', tab='usersTab'))
        
    # Защита: нельзя удалить самого себя
    if current_user.id in user_ids:
        user_ids.remove(current_user.id)
        
    if not user_ids:
        flash('Нельзя удалить собственную учетную запись', 'danger')
        return redirect(url_for('admin.dashboard', tab='usersTab'))

    # Защита для менеджера: только пользователи его отдела
    if current_user.role == 'manager':
        from app.utils import get_manager_department
        dept = get_manager_department(current_user)
        if not dept:
            flash('Доступ запрещен', 'danger')
            return redirect(url_for('admin.dashboard', tab='usersTab'))
        allowed_users = User.query.filter(User.id.in_(user_ids), User.department_id == dept.id, User.role == 'user').all()
        target_ids = [u.id for u in allowed_users]
    else:
        # Для админа: исключаем других пользователей с ролью admin во избежание случайного удаления
        target_users = User.query.filter(User.id.in_(user_ids), User.role != 'admin').all()
        target_ids = [u.id for u in target_users]
        
    if not target_ids:
        flash('Нет доступных для удаления пользователей (аккаунты администраторов защищены)', 'warning')
        return redirect(url_for('admin.dashboard', tab='usersTab'))
        
    from app.models import ReportDraft, ActionLog, UploadedFile
    from app.models.tasks import BackgroundTask
    
    # Каскадная очистка связанных записей
    ReportSubmission.query.filter(ReportSubmission.user_id.in_(target_ids)).delete(synchronize_session=False)
    ReportDraft.query.filter(ReportDraft.user_id.in_(target_ids)).delete(synchronize_session=False)
    ActionLog.query.filter(ActionLog.user_id.in_(target_ids)).update({'user_id': None}, synchronize_session=False)
    BackgroundTask.query.filter(BackgroundTask.user_id.in_(target_ids)).update({'user_id': None}, synchronize_session=False)
    UploadedFile.query.filter(UploadedFile.uploader_id.in_(target_ids)).update({'uploader_id': None}, synchronize_session=False)
    
    users_to_delete = User.query.filter(User.id.in_(target_ids)).all()
    deleted_count = len(users_to_delete)
    usernames_sample = ', '.join([u.username for u in users_to_delete[:5]])
    if deleted_count > 5:
        usernames_sample += f' и еще {deleted_count - 5}'
        
    for u in users_to_delete:
        u.assigned_templates = []
        db.session.delete(u)
        
    db.session.commit()
    log_action('Массовое удаление', f'Удалено пользователей: {deleted_count} ({usernames_sample})')
    flash(f'Успешно удалено пользователей: {deleted_count}', 'success')
    return redirect(url_for('admin.dashboard', tab='usersTab'))

